import io

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

@pytest.fixture
def client(tmp_path, monkeypatch):
    from backend import database
    monkeypatch.setattr(database, "DATA_DIR", tmp_path)
    from backend import main
    monkeypatch.setattr(main, "DATA_DIR", tmp_path)
    with TestClient(main.app) as c:
        yield c

def create_candidate(client):
    p = client.post("/api/projects", json={"title": "本地面试", "description": "开发系统", "kind": "project"})
    assert p.status_code == 201
    c = client.post("/api/candidates", json={"project_id": p.json()["id"], "name": "张三", "role": "开发", "notes": "备注"})
    assert c.status_code == 201
    return c.json()["id"]

def test_formats_and_persistence(client):
    from backend import main, database
    candidate_id = create_candidate(client)
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Resume Python FastAPI")
    pdf_bytes = pdf.tobytes()
    pdf.close()
    doc = Document()
    doc.add_paragraph("简历 Python FastAPI")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "表格技能 React"
    stream = io.BytesIO()
    doc.save(stream)
    samples = [("resume.pdf", pdf_bytes, "Python FastAPI"), ("resume.docx", stream.getvalue(), "表格技能 React"), ("resume.txt", "张三 Python 开发".encode(), "张三 Python"), ("gbk.txt", "中文简历".encode("gb18030"), "中文简历"), ("utf16.txt", "中文简历".encode("utf-16"), "中文简历")]
    for name, content, expected in samples:
        result = client.post(f"/api/candidates/{candidate_id}/resumes", files={"file": (name, content)})
        assert result.status_code == 201, result.text
        resume = result.json()
        assert expected in resume["text"]
        assert (database.DATA_DIR / resume["file_path"]).exists()
        assert client.get(f"/api/resumes/{resume['id']}").json()["text"] == resume["text"]
    # Repeat migration without losing data; process restart has a separate test.
    database.initialize()
    assert len(client.get("/api/projects").json()) == 1
    assert client.get("/api/projects/1/candidates").json()[0]["name"] == "张三"
    assert len(client.get(f"/api/candidates/{candidate_id}/resumes").json()) == len(samples)

def test_upload_failures_cleanup_and_empty_pdf(client):
    from backend import database
    candidate_id = create_candidate(client)
    for name, data, status in [("bad.exe", b"test", 400), ("empty.txt", b"", 400), ("bad.pdf", b"not a PDF", 422), ("bad.docx", b"not a DOCX", 422), ("large.txt", b"x" * (20 * 1024 * 1024 + 1), 413)]:
        assert client.post(f"/api/candidates/{candidate_id}/resumes", files={"file": (name, data)}).status_code == status
        assert list((database.DATA_DIR / "resumes").iterdir()) == []
    assert client.post("/api/candidates/999/resumes", files={"file": ("resume.txt", b"text")}).status_code == 404
    pdf = pymupdf.open()
    pdf.new_page()
    result = client.post(f"/api/candidates/{candidate_id}/resumes", files={"file": ("scan.pdf", pdf.tobytes())})
    pdf.close()
    assert result.status_code == 201
    assert result.json()["warning"] and result.json()["text"] == ""

def test_validation_and_settings_secret(client):
    assert client.post("/api/projects", json={"title": "  ", "description": "x"}).status_code == 422
    assert client.post("/api/candidates", json={"project_id": 999, "name": "张三", "role": "开发"}).status_code == 404
    payload = {"provider": "custom", "base_url": "http://localhost:11434/v1", "model": "local-model", "api_key": "test-secret-not-real"}
    response = client.put("/api/settings", json=payload)
    assert response.status_code == 200 and response.json()["has_api_key"]
    assert "test-secret-not-real" not in response.text
    assert "api_key" not in client.get("/api/settings").json()
    payload["api_key"] = None
    assert client.put("/api/settings", json=payload).json()["has_api_key"]
    payload["api_key"] = ""
    assert not client.put("/api/settings", json=payload).json()["has_api_key"]

def test_batch_item_import_is_atomic_and_separate(client):
    from backend import database
    project=client.post('/api/projects',json={'title':'批量导入','description':'开发系统'}).json()
    route=f"/api/projects/{project['id']}/import-candidate"
    first=client.post(route,data={'name':'林初','role':'后端'},files={'file':('lin.txt','林初的简历'.encode())})
    assert first.status_code==201,first.text
    second=client.post(route,data={'name':'周宁','role':'后端'},files={'file':('zhou.txt','周宁的简历'.encode())})
    assert second.status_code==201
    assert first.json()['candidate']['id']!=second.json()['candidate']['id']
    for name,data,status in [('broken.pdf',b'bad-pdf',422),('bad.exe',b'x',400),('empty.txt',b'',400)]:
        assert client.post(route,data={'name':'失败项','role':'后端'},files={'file':(name,data)}).status_code==status
    assert client.post(route,data={'name':'  ','role':'后端'},files={'file':('a.txt',b'abc')}).status_code==422
    people=client.get(f"/api/projects/{project['id']}/candidates").json()
    assert len(people)==2 and all(p['resume_count']==1 and p['session_count']==0 for p in people)
    assert len(list((database.DATA_DIR/'resumes').iterdir()))==2
    for person in (first.json(),second.json()):
        resumes=client.get(f"/api/candidates/{person['candidate']['id']}/resumes").json()
        assert len(resumes)==1 and resumes[0]['id']==person['resume']['id']
