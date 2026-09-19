from pathlib import Path
from zipfile import ZipFile
import pymupdf
from docx import Document

SUPPORTED = {".pdf", ".docx", ".txt"}
MAX_UPLOAD = 20 * 1024 * 1024
MAX_TEXT = 2_000_000

def parse_resume(path: Path) -> tuple[str, str]:
    try:
        if path.suffix == ".pdf":
            with pymupdf.open(stream=path.read_bytes(), filetype="pdf") as doc:
                if doc.needs_pass:
                    raise ValueError("PDF 已加密，请先移除密码。")
                if len(doc) > 200:
                    raise ValueError("PDF 不得超过 200 页。")
                text = "\n".join(page.get_text() for page in doc)
        elif path.suffix == ".docx":
            with ZipFile(path) as archive:
                if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("DOCX 解压内容过大。")
            doc = Document(path)
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                parts.extend("\t".join(c.text for c in row.cells) for row in table.rows)
            text = "\n".join(parts)
        else:
            data = path.read_bytes()
            encodings = ["utf-16"] if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "gb18030"]
            for encoding in encodings:
                try:
                    text = data.decode(encoding)
                    break
                except UnicodeError:
                    continue
            else:
                raise ValueError("TXT 编码无法识别，请另存为 UTF-8。")
            if "\x00" in text:
                raise ValueError("TXT 含二进制内容，请上传文本文件。")
        if len(text) > MAX_TEXT:
            raise ValueError("解析文本过长，请精简简历后重试。")
        text = text.strip()
        warning = "" if text else "未提取到文本。扫描 PDF 需要 OCR，第一阶段暂不支持；请上传含文字的 PDF、DOCX 或 TXT。"
        return text, warning
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("无法解析文件，请确认格式正确、文件未损坏且未加密。") from exc
